from docx import Document
from docx.shared import Pt
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT


def generate_falc_comparison_doc(examples, output_file="Comparaison_FALC.docx"):
    """
    examples = [
        {
            "original": "...",
            "generated": "...",
            "emotion": "Tristesse",
            "emotion_score": 0.91
        },
        ...
    ]
    """

    doc = Document()

    title = doc.add_heading("Exemples de génération FALC", level=1)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    intro = doc.add_paragraph(
        "Les exemples suivants présentent le texte original et la version "
        "FALC générée par le modèle."
    )
    intro.style.font.size = Pt(11)

    for i, example in enumerate(examples, start=1):

        doc.add_heading(f"Exemple {i}", level=2)

        # Optional emotion information
        if "emotion" in example:
            p = doc.add_paragraph()
            p.add_run("Émotion principale : ").bold = True
            p.add_run(example["emotion"])

            if "emotion_score" in example:
                p.add_run(
                    f" (score : {example['emotion_score']:.2f})"
                )

        table = doc.add_table(rows=2, cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Header
        table.cell(0, 0).text = "Texte original"
        table.cell(0, 1).text = "Texte FALC généré"

        # Content
        table.cell(1, 0).text = example["original"]
        table.cell(1, 1).text = example["generated"]

        doc.add_paragraph()

    doc.save(output_file)
    print(f"Saved to {output_file}")






examples = [
    {
        "original":
        """Léa, pourquoi es-tu triste ? Ce matin, Clément se réveille tout content.
Aujourd'hui, c'est l'anniversaire de sa petite sœur Léa...""",

        "generated":
        """Léa semble triste aujourd'hui.
Clément est heureux.
C'est l'anniversaire de Léa.
Il court dans sa chambre pour lui dire :
« Bon anniversaire ! »
Léa est très fatiguée.
Elle a mal dormi cette nuit.
Elle voulait encore dormir un peu.""",

        "emotion": "Tristesse",
        "emotion_score": 0.93
    },

    {
        "original":
        "Paul est très heureux de retrouver ses amis.",

        "generated":
        "Paul est très content. Il est heureux de revoir ses amis.",

        "emotion": "Joie",
        "emotion_score": 0.97
    }
]

generate_falc_comparison_doc(examples)